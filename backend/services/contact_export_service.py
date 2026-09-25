"""
Экспорт базы контактов агента обзвона в Excel (.xlsx).

Один файл, два листа:
  - «Контакты» — одна строка на контакт: данные контакта, стадия воронки,
    итог последнего звонка, память агента (итог, факты, лучшее время),
    ближайший запланированный шаг. Первые пять колонок совпадают с шаблоном
    импорта (generate_template_xlsx), поэтому файл можно отредактировать и
    загрузить обратно через «Импорт».
  - «Звонки» — одна строка на каждый завершённый разговор/сообщение:
    контакт, дата, канал, направление, статус, решение агента, длительность,
    транскрипт.

Время везде в МСК, как в интерфейсе. Файл собирается в памяти (openpyxl).
"""

import io
from datetime import datetime
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from backend.core.pipeline_stages import AGENT_CONTACT_STAGES
from backend.core.timezone_utils import utc_to_msk
from backend.models.agent_call import AgentCall
from backend.models.agent_contact import AgentContact
from backend.models.task import Task, TaskStatus

# Лимит длины текстовой ячейки Excel (32 767 символов); оставляем запас под пометку.
_CELL_MAX = 32_000
_TRUNCATED_MARK = "\n… [обрезано]"

_STAGE_LABELS = {s["key"]: s["label"] for s in AGENT_CONTACT_STAGES}

_DECISION_LABELS = {
    "SUCCESS": "Успех",
    "FOLLOWUP": "Перезвонить",
    "NO_ANSWER": "Не дозвонились",
    "REJECTED": "Отказ",
    "DO_NOT_CALL": "Не звонить",
}

_CALL_STATUS_LABELS = {
    "answered": "Ответил",
    "no_answer": "Не ответил",
    "failed": "Ошибка",
}

_CHANNEL_LABELS = {
    "call": "Звонок",
    "sms": "SMS",
    "telegram": "Telegram",
    "max": "MAX",
}

_DIRECTION_LABELS = {
    "outbound": "Исходящий",
    "inbound": "Входящий",
}

# Только завершённые события — как в списке звонков UI (/api/agent/calls).
FINALIZED_CALL_STATUSES = ["answered", "no_answer", "failed"]

CONTACT_HEADERS = [
    "Имя", "Телефон", "Компания", "Должность", "Информация о клиенте",
    "Стадия", "Попыток", "Последний звонок (МСК)", "Результат последнего звонка",
    "Итог общения", "Ключевые факты", "Лучшее время для звонка",
    "Следующий шаг", "Добавлен (МСК)",
]
CONTACT_WIDTHS = [20, 16, 20, 16, 36, 12, 9, 18, 22, 44, 44, 20, 40, 18]

CALL_HEADERS = [
    "Контакт", "Телефон", "Дата (МСК)", "Канал", "Направление",
    "Статус", "Решение агента", "Длительность (сек)", "Транскрипт",
    "Запись звонка",
]
CALL_WIDTHS = [20, 16, 18, 10, 12, 12, 18, 14, 80, 40]


def _fmt_dt(dt: Optional[datetime]) -> str:
    if not dt:
        return ""
    return utc_to_msk(dt).strftime("%d.%m.%Y %H:%M")


def _cell(value: Any) -> Any:
    """
    Значение для ячейки: None → "", строки обрезаем до лимита Excel и
    защищаем от трактовки как формулы (openpyxl считает строку с '=' формулой).
    """
    if value is None:
        return ""
    if isinstance(value, (int, float)):
        return value
    s = str(value)
    if len(s) > _CELL_MAX:
        s = s[:_CELL_MAX] + _TRUNCATED_MARK
    if s.startswith("="):
        s = " " + s
    return s


def _join_list(items: Any) -> str:
    if not items:
        return ""
    if isinstance(items, str):
        return items
    try:
        return "; ".join(str(i) for i in items if i)
    except TypeError:
        return str(items)


def _next_step(task: Optional[Task]) -> str:
    if not task:
        return ""
    when = _fmt_dt(task.scheduled_time)
    channel = _CHANNEL_LABELS.get(task.channel or "call", task.channel or "")
    text = task.description or task.title or ""
    parts = [p for p in (when, channel, text) if p]
    return " · ".join(parts)


def _style_header(ws, widths: List[int]) -> None:
    fill = PatternFill("solid", fgColor="E8EEF9")
    for i, w in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _collect(db: Session, agent_config_id) -> Dict[str, Any]:
    """Достаём всё одним проходом: контакты, звонки, ближайшие задачи."""
    contacts: List[AgentContact] = (
        db.query(AgentContact)
        .filter(AgentContact.agent_config_id == agent_config_id)
        .order_by(AgentContact.created_at.desc())
        .all()
    )

    calls: List[AgentCall] = (
        db.query(AgentCall)
        .filter(
            AgentCall.agent_config_id == agent_config_id,
            AgentCall.status.in_(FINALIZED_CALL_STATUSES),
        )
        .order_by(AgentCall.created_at.desc())
        .all()
    )

    # Последний завершённый звонок по каждому контакту (список уже отсортирован desc).
    last_call_by_contact: Dict[str, AgentCall] = {}
    for c in calls:
        if c.agent_contact_id and str(c.agent_contact_id) not in last_call_by_contact:
            last_call_by_contact[str(c.agent_contact_id)] = c

    # Ближайшая запланированная задача по каждому контакту.
    contact_ids = [c.id for c in contacts]
    next_task_by_contact: Dict[str, Task] = {}
    if contact_ids:
        tasks = (
            db.query(Task)
            .filter(
                Task.agent_contact_id.in_(contact_ids),
                Task.is_agent_task == True,  # noqa: E712
                Task.status == TaskStatus.SCHEDULED,
            )
            .order_by(Task.scheduled_time.asc())
            .all()
        )
        for t in tasks:
            key = str(t.agent_contact_id)
            if key not in next_task_by_contact:
                next_task_by_contact[key] = t

    return {
        "contacts": contacts,
        "calls": calls,
        "last_call_by_contact": last_call_by_contact,
        "next_task_by_contact": next_task_by_contact,
    }


def generate_contacts_export_xlsx(db: Session, agent_config_id) -> bytes:
    """Собирает xlsx со всей базой контактов агента и историей звонков."""
    data = _collect(db, agent_config_id)
    contacts: List[AgentContact] = data["contacts"]
    calls: List[AgentCall] = data["calls"]
    last_call_by_contact = data["last_call_by_contact"]
    next_task_by_contact = data["next_task_by_contact"]

    wb = Workbook()

    # ── Лист 1: Контакты ──
    ws = wb.active
    ws.title = "Контакты"
    ws.append(CONTACT_HEADERS)
    for c in contacts:
        mem = c.memory if isinstance(c.memory, dict) else {}
        last_call = last_call_by_contact.get(str(c.id))
        decision = last_call.post_call_decision if last_call else None
        ws.append([
            _cell(c.name),
            _cell(c.phone),
            _cell(c.company),
            _cell(c.position),
            _cell(c.notes),
            # Легаси-статусы (напр. "calling") в UI показываются как «В работе».
            _cell(_STAGE_LABELS.get(c.status) or _STAGE_LABELS["active"]),
            int(c.attempts_count or 0),
            _cell(_fmt_dt(c.last_called_at)),
            _cell(_DECISION_LABELS.get(decision, decision) if decision else ""),
            _cell(mem.get("summary")),
            _cell(_join_list(mem.get("key_facts"))),
            _cell(mem.get("best_time")),
            _cell(_next_step(next_task_by_contact.get(str(c.id)))),
            _cell(_fmt_dt(c.created_at)),
        ])
    _style_header(ws, CONTACT_WIDTHS)

    # ── Лист 2: Звонки ──
    ws2 = wb.create_sheet("Звонки")
    ws2.append(CALL_HEADERS)
    contact_by_id = {str(c.id): c for c in contacts}
    for call in calls:
        contact = contact_by_id.get(str(call.agent_contact_id)) if call.agent_contact_id else None
        # Для истории берём момент начала, иначе создание записи.
        when = call.started_at or call.created_at
        channel = call._resolve_channel()
        decision = call.post_call_decision
        ws2.append([
            _cell(contact.name if contact else ""),
            _cell(contact.phone if contact else ""),
            _cell(_fmt_dt(when)),
            _cell(_CHANNEL_LABELS.get(channel, channel)),
            _cell(_DIRECTION_LABELS.get(call.direction or "outbound", call.direction)),
            _cell(_CALL_STATUS_LABELS.get(call.status, call.status)),
            _cell(_DECISION_LABELS.get(decision, decision) if decision else ""),
            int(call.duration_seconds or 0),
            _cell(call.transcript),
            _cell(call.record_url or ""),
        ])
    _style_header(ws2, CALL_WIDTHS)
    for row in ws2.iter_rows(min_row=2, min_col=9, max_col=9):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
